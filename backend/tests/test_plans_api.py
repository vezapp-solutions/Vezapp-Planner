"""Plans API: create and recalculate, line detail, load, bottlenecks, Gantt, lifecycle, roles, tenancy."""

import random
import time
from datetime import date, timedelta

import pytest

from apps.masters.models import Customer, Part
from apps.plans.models import Plan, PlanAllocation
from apps.scheduler.models import ScheduleLine
from seeds.demo_a import build_demo_foundry
from seeds.worked_example import build_worked_example

J = {"format": "json"}
TODAY = date(2026, 11, 2)  # a Monday


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr("django.utils.timezone.localdate", lambda *a, **k: TODAY)


@pytest.fixture
def example(company_a):
    """Company A holds the section 7.5 plant and one schedule line for part H-120, 600 pieces."""
    built = build_worked_example(company_a)
    customer = Customer.objects.get(company=company_a, code="CUST-01")
    ScheduleLine.objects.create(
        company=company_a, customer=customer, part=built["part"], dispatch_qty=600,
        internal_dispatch_date=date(2026, 11, 20),
    )  # fmt: skip
    return built


def rows(response):
    return response.json()["results"]


def make_plan(client):
    response = client.post("/api/v1/plans", {}, **J)
    assert response.status_code == 201, response.content
    return response.json()


# --- the section 7.5 example through the whole stack -------------------------------------------------------
def test_planner_creates_a_plan_and_the_worked_example_numbers_come_back(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    assert plan["version"] == 1 and plan["status"] == "draft" and plan["calculated_at"]
    assert plan["bottleneck"]["code"] == "F1"
    assert plan["summary"]["by_status"] == {"on_track": 1}
    assert not plan["stale"]

    line = rows(client.get(f"/api/v1/plans/{plan['id']}/lines"))[0]
    assert (line["part_number"], line["status"], line["est_dispatch_date"]) == (
        "H-120",
        "on_track",
        "2026-11-05",
    )

    detail = client.get(f"/api/v1/plans/{plan['id']}/lines/{line['id']}").json()
    req = detail["requirement"]
    assert (req["pieces_to_pour"], req["moulds"], req["cores"], req["metal_kg"], req["heats"]) == (
        630,
        315,
        630,
        18900.0,
        21,
    )
    assert req["minutes"] == {
        "core": 630.0,
        "moulding": 472.5,
        "melting": 1575.0,
        "knockout": 315.0,
        "fettling": 1890.0,
    }
    assert req["traces"]["heats"]["formula"].startswith("ceil(moulds")

    moulding = {}
    for a in detail["allocations"]:
        if a["op_code"] == "moulding":
            moulding[a["date"]] = moulding.get(a["date"], 0) + a["qty"]
    assert list(moulding.values()) == [150.0, 150.0, 15.0]  # three days: the furnace paces the line
    assert [o["code"] for o in detail["operations"]] == [
        "core",
        "moulding",
        "melting",
        "knockout",
        "fettling",
    ]


def test_load_grid_shows_the_furnace_at_98_percent_on_the_first_two_days(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    grid = client.get(f"/api/v1/plans/{plan['id']}/load?from=2026-11-02&to=2026-11-06").json()
    furnace = next(r for r in grid["resources"] if r["code"] == "F1")
    day1 = furnace["cells"]["2026-11-02"]
    assert (
        day1["load"] == 750.0
        and day1["capacity"] == 765.0
        and day1["utilisation"] == pytest.approx(0.9804, abs=1e-4)
    )
    assert furnace["cells"]["2026-11-04"]["load"] == 75.0
    by_shift = client.get(
        f"/api/v1/plans/{plan['id']}/load?from=2026-11-02&to=2026-11-02&by=shift"
    ).json()
    assert [c["shift"] for c in by_shift["columns"]] == ["Shift 1", "Shift 2"]
    f_shift = next(r for r in by_shift["resources"] if r["code"] == "F1")
    assert [c["load"] for c in f_shift["cells"].values()] == [375.0, 375.0]


def test_load_cell_lists_the_lines_behind_a_number(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    furnace = example["furnace"]
    cell = client.get(
        f"/api/v1/plans/{plan['id']}/load/cell?resource={furnace.pk}&date=2026-11-02"
    ).json()
    assert [(r["part"], r["op_code"], r["minutes"]) for r in cell["lines"]] == [
        ("H-120", "melting", 750.0)
    ]


def test_bottleneck_strip_and_day_view(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    strip = client.get(f"/api/v1/plans/{plan['id']}/bottlenecks?days=14").json()
    assert len(strip) == 14 and strip[0]["date"] == "2026-11-02"
    assert strip[0]["resource_code"] == "F1" and strip[1]["resource_code"] == "F1"
    assert strip[6]["resource_code"] is None  # nothing planned that far out
    assert "2026-11-08" not in [s["date"] for s in strip]  # Sunday is a weekly off

    day = client.get(f"/api/v1/plans/{plan['id']}/day?date=2026-11-02").json()
    assert {w["op_code"]: w["qty"] for w in day["work"]}["moulding"] == 150.0
    assert next(r for r in day["resources"] if r["code"] == "F1")["utilisation"] == pytest.approx(
        0.9804, abs=1e-4
    )


def test_gantt_by_resource_and_by_line(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    by_resource = client.get(
        f"/api/v1/plans/{plan['id']}/gantt?from=2026-11-02&to=2026-11-08"
    ).json()
    assert {b["row_label"] for b in by_resource["bars"]} == {"CM1", "LINE-A", "F1", "KO", "FT"}
    by_line = client.get(
        f"/api/v1/plans/{plan['id']}/gantt?by=line&from=2026-11-02&to=2026-11-08"
    ).json()
    assert [b["code"] for b in by_line["bars"]] == [
        "core",
        "moulding",
        "melting",
        "knockout",
        "fettling",
    ]
    assert (
        client.get(f"/api/v1/plans/{plan['id']}/gantt?from=2026-01-01&to=2026-12-31").status_code
        == 400
    )


def test_lines_are_sorted_most_urgent_first_and_filterable(login, users, example, company_a):
    customer = Customer.objects.get(company=company_a, code="CUST-01")
    ScheduleLine.objects.create(
        company=company_a,
        customer=customer,
        part=example["part"],
        dispatch_qty=600,
        internal_dispatch_date=date(2026, 11, 3),
    )
    client = login(users["a_planner"])
    plan = make_plan(client)
    got = rows(client.get(f"/api/v1/plans/{plan['id']}/lines"))
    assert [r["status"] for r in got][0] == "at_risk"
    assert len(rows(client.get(f"/api/v1/plans/{plan['id']}/lines?status=at_risk"))) == 1
    assert len(rows(client.get(f"/api/v1/plans/{plan['id']}/lines?month=2026-11"))) == 2
    assert rows(client.get(f"/api/v1/plans/{plan['id']}/lines?q=zzz")) == []
    assert client.get(f"/api/v1/plans/{plan['id']}/lines?month=bad").status_code == 400


# --- lifecycle ---------------------------------------------------------------------------------------------------
def test_changes_make_the_plan_stale_and_recalculating_clears_it(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    ScheduleLine.objects.filter(company=example["part"].company).update(dispatch_qty=700)
    ScheduleLine.objects.first().save()  # touches updated_at, as an edit through the API would
    stale = client.get(f"/api/v1/plans/{plan['id']}").json()
    assert stale["stale"] and "changed" in stale["stale_reasons"][0]
    fresh = client.post(f"/api/v1/plans/{plan['id']}/recalculate", {}, **J).json()
    assert not fresh["stale"]
    part = Part.objects.get(pk=example["part"].pk)
    part.casting_weight_kg = 26
    part.save()
    assert client.get(f"/api/v1/plans/{plan['id']}").json()["stale"]  # masters count too


def test_a_plan_from_an_earlier_day_is_flagged(login, users, example, monkeypatch):
    client = login(users["a_planner"])
    plan = make_plan(client)
    monkeypatch.setattr(
        "django.utils.timezone.localdate", lambda *a, **k: TODAY + timedelta(days=1)
    )
    got = client.get(f"/api/v1/plans/{plan['id']}").json()
    assert got["stale"] and "earlier day" in got["stale_reasons"][-1]


def test_confirm_freezes_a_plan_and_needs_it_to_be_current(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    ScheduleLine.objects.first().save()
    refused = client.post(f"/api/v1/plans/{plan['id']}/confirm", {}, **J)
    assert refused.status_code == 400 and "out of date" in refused.json()["detail"]

    client.post(f"/api/v1/plans/{plan['id']}/recalculate", {}, **J)
    confirmed = client.post(f"/api/v1/plans/{plan['id']}/confirm", {}, **J).json()
    assert confirmed["status"] == "confirmed" and confirmed["confirmed_at"]
    assert client.post(f"/api/v1/plans/{plan['id']}/recalculate", {}, **J).status_code == 400
    assert client.post(f"/api/v1/plans/{plan['id']}/confirm", {}, **J).status_code == 400
    assert client.delete(f"/api/v1/plans/{plan['id']}").status_code == 400
    ScheduleLine.objects.first().save()  # a baseline does not go stale
    assert not client.get(f"/api/v1/plans/{plan['id']}").json()["stale"]


def test_versions_increase_and_latest_prefers_the_newest_draft(login, users, example):
    client = login(users["a_planner"])
    assert client.get("/api/v1/plans/latest").status_code == 204  # nothing planned yet
    first = make_plan(client)
    client.post(f"/api/v1/plans/{first['id']}/confirm", {}, **J)
    second = make_plan(client)
    assert (first["version"], second["version"]) == (1, 2)
    assert client.get("/api/v1/plans/latest").json()["id"] == second["id"]
    assert client.delete(f"/api/v1/plans/{second['id']}").status_code == 204
    assert client.get("/api/v1/plans/latest").json()["id"] == first["id"]


def test_recalculating_replaces_the_results_instead_of_adding_to_them(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    before = PlanAllocation.objects.filter(plan_id=plan["id"]).count()
    client.post(f"/api/v1/plans/{plan['id']}/recalculate", {}, **J)
    assert PlanAllocation.objects.filter(plan_id=plan["id"]).count() == before


def test_a_company_with_no_lines_gets_an_empty_plan(login, users, company_a):
    plan = make_plan(login(users["a_planner"]))
    assert plan["summary"]["lines"] == 0 and plan["bottleneck"] is None


# --- roles and tenancy --------------------------------------------------------------------------------------------
@pytest.mark.parametrize("who, expected", [("a_admin", 201), ("a_planner", 201), ("a_viewer", 403)])
def test_roles_for_creating_plans(login, users, example, who, expected):
    assert login(users[who]).post("/api/v1/plans", {}, **J).status_code == expected


def test_viewers_can_read_plans_but_not_change_them(login, users, example):
    plan = make_plan(login(users["a_planner"]))
    viewer = login(users["a_viewer"])
    for path in ("", "/lines", "/load", "/bottlenecks", "/gantt", "/day"):
        assert viewer.get(f"/api/v1/plans/{plan['id']}{path}").status_code == 200, path
    for action in ("recalculate", "confirm"):
        assert viewer.post(f"/api/v1/plans/{plan['id']}/{action}", {}, **J).status_code == 403
    assert viewer.delete(f"/api/v1/plans/{plan['id']}").status_code == 403


def test_company_a_never_sees_company_bs_plans(login, users, example, company_b):
    b_part = build_worked_example(company_b)["part"]
    ScheduleLine.objects.create(
        company=company_b, customer=Customer.objects.get(company=company_b, code="CUST-01"),
        part=b_part, dispatch_qty=100, internal_dispatch_date=date(2026, 11, 20),
    )  # fmt: skip
    theirs = make_plan(login(users["b_admin"]))
    mine = make_plan(login(users["a_admin"]))
    client = login(users["a_admin"])
    assert [p["id"] for p in rows(client.get("/api/v1/plans"))] == [mine["id"]]
    for path in ("", "/lines", "/load", "/bottlenecks", "/gantt", "/day", "/lines/1"):
        assert client.get(f"/api/v1/plans/{theirs['id']}{path}").status_code == 404, path
    assert client.post(f"/api/v1/plans/{theirs['id']}/recalculate", {}, **J).status_code == 404
    assert client.delete(f"/api/v1/plans/{theirs['id']}").status_code == 404
    assert Plan.objects.get(pk=theirs["id"]).status == "draft"
    assert client.get("/api/v1/plans/latest").json()["id"] == mine["id"]


def test_vezapp_admin_plans_for_a_company_with_the_header(login, users, example, company_a):
    client = login(users["vezapp_admin"])
    assert client.post("/api/v1/plans", {}, **J).status_code == 400  # no company chosen
    client.credentials(HTTP_X_VP_CLIENT="web", HTTP_X_COMPANY_ID=str(company_a.pk))
    assert client.post("/api/v1/plans", {}, **J).status_code == 201


def test_plan_creation_is_audited(login, users, example):
    from apps.core.models import AuditLog

    make_plan(login(users["a_planner"]))
    assert AuditLog.objects.filter(
        model_label="plans.Plan", action="create", username="a_planner"
    ).exists()


# --- a realistic month through the API, and the speed target ----------------------------------------------------------
def test_replanning_500_lines_through_the_api_takes_under_five_seconds(login, users, company_a):
    build_demo_foundry(company_a)
    rng = random.Random(3)
    pairs = [
        (link.customer, p)
        for p in Part.objects.filter(company=company_a)
        for link in p.customer_links.all()
    ]
    ScheduleLine.objects.bulk_create(
        [
            ScheduleLine(
                company=company_a, customer=c, part=p, dispatch_qty=rng.choice([20, 40, 60, 100, 150]),
                internal_dispatch_date=TODAY + timedelta(days=rng.randint(3, 120)),
            )
            for c, p in (rng.choice(pairs) for _ in range(500))
        ]
    )  # fmt: skip
    client = login(users["a_planner"])
    started = time.perf_counter()
    plan = make_plan(client)
    elapsed = time.perf_counter() - started
    assert plan["summary"]["lines"] == 500
    assert elapsed < 5, f"creating a plan for 500 lines took {elapsed:.1f}s"
    started = time.perf_counter()
    assert client.post(f"/api/v1/plans/{plan['id']}/recalculate", {}, **J).status_code == 200
    assert time.perf_counter() - started < 5


# --- printed outputs ---------------------------------------------------------------------------------------------
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sheet_text(content):
    import io

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content)).active
    return [[c.value for c in row] for row in sheet.iter_rows()]


def pdf_pages(content):
    import io

    from pypdf import PdfReader

    return PdfReader(io.BytesIO(content)).pages


def test_shift_plan_pdf_has_the_logo_title_and_the_days_work(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=shift_plan&department=moulding&date=2026-11-02&output=pdf"
    )
    assert response.status_code == 200 and response["Content-Type"] == "application/pdf"
    assert (
        response.content.startswith(b"%PDF-")
        and "shift-plan-2026-11-02-moulding.pdf" in response["Content-Disposition"]
    )
    pages = pdf_pages(response.content)
    text = pages[0].extract_text()
    assert (
        "Shift plan: Moulding" in text
        and "02-11-2026" in text
        and "Shift 1" in text
        and "Shift 2" in text
    )
    assert "H-120" in text and "75 moulds" in text and "LINE-A" in text
    assert "Utilisation" in text and len(pages[0].images) >= 1  # the Vezapp logo is on the page


def test_shift_plan_excel_matches_and_stores_text_only(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=shift_plan&department=fettling&date=2026-11-03&shift=1&output=xlsx"
    )
    assert response.status_code == 200 and response["Content-Type"] == XLSX
    flat = [str(c) for row in sheet_text(response.content) for c in row if c is not None]
    assert "Shift plan: Fettling" in flat and "H-120" in flat and "FT" in flat
    assert not any(v.startswith("=") for v in flat)


def test_day_plan_adds_the_shifts_together(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=day_plan&department=moulding&date=2026-11-02&output=xlsx"
    )
    flat = [str(c) for row in sheet_text(response.content) for c in row if c is not None]
    assert "150 moulds" in flat  # 75 + 75 over both shifts


def test_melting_plan_is_in_moulds_not_heats(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=shift_plan&department=melting&date=2026-11-02&output=pdf"
    )
    text = pdf_pages(response.content)[0].extract_text()
    assert "F1" in text and "75 moulds" in text and "heat" not in text.lower()


def test_dispatch_vs_estimate_lists_each_line_with_dates_and_slack(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=dispatch_vs_estimate&month=2026-11&output=xlsx"
    )
    table = sheet_text(response.content)
    row = next(r for r in table if r and r[0] == "20-11-2026")
    assert row[1:5] == ["CUST-01", "H-120 Housing", "600", "05-11-2026"] and row[6] == "On track"
    assert int(row[5]) > 0
    summary = next(r for r in table if r and r[0] == "Lines")
    assert table[table.index(summary) + 1][:2] == ["1", "1"]


def test_resource_load_report_shows_utilisation_per_day(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=resource_load&from=2026-11-02&to=2026-11-06&output=xlsx"
    )
    table = sheet_text(response.content)
    header = next(r for r in table if r and r[0] == "Resource")
    furnace = next(r for r in table if r and r[0] == "F1")
    assert header[1:3] == ["02-11", "03-11"] and furnace[1] == "98%"


@pytest.mark.parametrize(
    "query",
    [
        "report=nonsense",
        "output=docx",
        "report=shift_plan&department=welding",
        "report=shift_plan&date=junk",
        "report=dispatch_vs_estimate&month=junk",
        "report=resource_load&from=2026-11-02&to=2026-01-01",
    ],
)
def test_bad_export_requests_are_refused_with_a_message(login, users, example, query):
    client = login(users["a_planner"])
    plan = make_plan(client)
    assert client.get(f"/api/v1/plans/{plan['id']}/export?{query}").status_code == 400


def test_a_day_with_no_work_prints_a_clear_note(login, users, example):
    client = login(users["a_planner"])
    plan = make_plan(client)
    response = client.get(
        f"/api/v1/plans/{plan['id']}/export?report=shift_plan&department=core&date=2026-12-20&output=xlsx"
    )
    flat = [str(c) for row in sheet_text(response.content) for c in row if c is not None]
    assert "No work is planned for this selection." in flat


def test_viewers_can_print_but_other_companies_plans_stay_out_of_reach(
    login, users, example, company_b
):
    plan = make_plan(login(users["a_planner"]))
    assert (
        login(users["a_viewer"]).get(f"/api/v1/plans/{plan['id']}/export?output=pdf").status_code
        == 200
    )
    assert (
        login(users["b_admin"]).get(f"/api/v1/plans/{plan['id']}/export?output=pdf").status_code
        == 404
    )
