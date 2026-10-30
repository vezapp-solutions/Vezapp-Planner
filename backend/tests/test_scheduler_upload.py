"""The Phase 2 acceptance case: a 200-row Excel file with deliberate errors is validated correctly."""

import io
import time
from datetime import date, datetime, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.core.models import AuditLog
from apps.masters.models import Customer, Grade, Part, PartCustomer
from apps.scheduler.importer import COLUMNS
from apps.scheduler.models import ScheduleLine, UploadBatch

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = [h for _, h in COLUMNS]
KEYS = [k for k, _ in COLUMNS]
TODAY = date.today()


@pytest.fixture
def world(company_a):
    """3 customers + 1 inactive, 5 active parts + 1 inactive, with a realistic link pattern."""
    grade = Grade.objects.create(company=company_a, name="GG25")
    customers = {
        code: Customer.objects.create(company=company_a, code=code, name=name, is_active=active)
        for code, name, active in (
            ("C1", "Aarav Pumps", True),
            ("C2", "Bharat Valves", True),
            ("C3", "Kaveri Auto", True),
            ("C4", "Dormant Ltd", False),
        )
    }
    parts = {
        n: Part.objects.create(
            company=company_a, part_number=n, name=f"Part {n}", grade=grade,
            casting_weight_kg=10, pouring_weight_per_mould_kg=25, cavities_per_mould=2,
            selling_price=1000 if n != "P5" else None, is_active=n != "P6",
        )
        for n in ("P1", "P2", "P3", "P4", "P5", "P6")
    }  # fmt: skip
    for c, plist in {"C1": "P1 P2 P3 P6", "C2": "P3 P4", "C3": "P4 P5", "C4": "P1"}.items():
        for p in plist.split():
            first = not PartCustomer.objects.filter(part=parts[p], is_default=True).exists()
            PartCustomer.objects.create(
                company=company_a, customer=customers[c], part=parts[p], is_default=first
            )
    return customers, parts


# valid, linked (customer, part) pairs: 3 + 2 + 2 = 7, times 28 days = 196 unique combinations
PAIRS = [
    ("C1", "P1"),
    ("C1", "P2"),
    ("C1", "P3"),
    ("C2", "P3"),
    ("C2", "P4"),
    ("C3", "P4"),
    ("C3", "P5"),
]
COMBOS = [(c, p, date(2026, 11, d)) for d in range(1, 29) for c, p in PAIRS]


def as_date_cell(day: date, style: int):
    """The same date as DD-MM-YYYY text, DD/MM/YYYY text, a real Excel date, or a serial number."""
    return [
        day.strftime("%d-%m-%Y"),
        day.strftime("%d/%m/%Y"),
        datetime(day.year, day.month, day.day),
        (day - date(1899, 12, 30)).days,
    ][style % 4]


def build_200_rows():
    """Returns [(values, expected_status, note)] in file order. Row 2 is the first entry."""
    rows = []

    def add(values, status, note):
        rows.append(({"customer_code": "", "customer_name": "", "part_number": "", "dispatch_qty": 100,
                      "internal_date": "", "actual_date": "", "actual_qty": "", "remarks": "", **values}, status, note))  # fmt: skip

    # 120 clean rows in every date and quantity style
    for i in range(120):
        c, p, d = COMBOS[i]
        qty = [500 + i, float(500 + i), f"{1000 + i:,}", str(700 + i)][i % 4]
        add(
            {
                "customer_code": c,
                "part_number": p,
                "dispatch_qty": qty,
                "internal_date": as_date_cell(d, i),
            },
            "ok",
            "clean",
        )

    # 11 clean partial dispatches (actual date in the past, actual quantity below the order)
    for i in range(120, 131):
        c, p, d = COMBOS[i]
        add({"customer_code": c, "part_number": p, "dispatch_qty": 400, "internal_date": as_date_cell(d, i),
             "actual_date": (TODAY - timedelta(days=3)).strftime("%d-%m-%Y"), "actual_qty": 250}, "ok", "partial")  # fmt: skip

    # 40 rows with errors
    for i in range(6):
        add(
            {"customer_code": f"NOPE{i}", "part_number": "P1", "internal_date": "10-11-2026"},
            "error",
            "unknown code",
        )
    for i in range(4):
        add(
            {
                "customer_name": f"No Such Customer {i}",
                "part_number": "P1",
                "internal_date": "10-11-2026",
            },
            "error",
            "unknown name",
        )
    for i in range(5):
        add(
            {"customer_code": "C1", "part_number": f"GHOST-{i}", "internal_date": "10-11-2026"},
            "error",
            "unknown part",
        )
    for bad in [0, -5, 12.5, "abc", ""]:
        add(
            {
                "customer_code": "C1",
                "part_number": "P1",
                "dispatch_qty": bad,
                "internal_date": "10-11-2026",
            },
            "error",
            "bad qty",
        )
    for bad in ["31-02-2026", "2026-13-45", "next week", "", 5]:
        add({"customer_code": "C1", "part_number": "P1", "internal_date": bad}, "error", "bad date")
    for i in range(4):
        future = (TODAY + timedelta(days=5 + i)).strftime("%d-%m-%Y")
        add(
            {
                "customer_code": "C1",
                "part_number": "P1",
                "internal_date": "10-11-2026",
                "actual_date": future,
            },
            "error",
            "future actual",
        )
    for _ in range(3):
        add(
            {
                "customer_code": "C1",
                "part_number": "P1",
                "internal_date": "10-11-2026",
                "actual_qty": 50,
            },
            "error",
            "actual qty without date",
        )
    for _ in range(2):
        add({"part_number": "P1", "internal_date": "10-11-2026"}, "error", "no customer")
    for i in range(6):
        add(
            {
                "customer_code": "C1",
                "part_number": f"GHOST-X{i}",
                "dispatch_qty": 0,
                "internal_date": "bad",
            },
            "error",
            "three problems",
        )

    # 24 warnings
    for i in range(8):  # part exists but is not linked to this customer
        add(
            {
                "customer_code": "C2",
                "part_number": "P1",
                "internal_date": date(2026, 11, 1 + i).strftime("%d-%m-%Y"),
            },
            "warning",
            "not linked",
        )
    for i in range(4):  # inactive part
        add(
            {
                "customer_code": "C1",
                "part_number": "P6",
                "internal_date": date(2026, 11, 1 + i).strftime("%d-%m-%Y"),
            },
            "warning",
            "inactive part",
        )
    for i in range(4):  # inactive customer
        add(
            {
                "customer_code": "C4",
                "part_number": "P1",
                "internal_date": date(2026, 11, 1 + i).strftime("%d-%m-%Y"),
            },
            "warning",
            "inactive customer",
        )
    for i in range(4):  # name does not match the code
        c, p, d = COMBOS[131 + i]
        add(
            {
                "customer_code": c,
                "customer_name": "Somebody Else",
                "part_number": p,
                "internal_date": d.strftime("%d-%m-%Y"),
            },
            "warning",
            "name differs",
        )
    for i in range(4):  # unknown code but exact name matches
        c, p, d = COMBOS[135 + i]
        name = {"C1": "Aarav Pumps", "C2": "Bharat Valves", "C3": "Kaveri Auto"}[c]
        add(
            {
                "customer_code": "OLD-CODE",
                "customer_name": name,
                "part_number": p,
                "internal_date": d.strftime("%d-%m-%Y"),
            },
            "warning",
            "code unknown",
        )

    # 5 duplicates of earlier clean rows (the later row is the one flagged)
    for i in range(5):
        c, p, d = COMBOS[i]
        add(
            {
                "customer_code": c,
                "part_number": p,
                "dispatch_qty": 123,
                "internal_date": d.strftime("%d-%m-%Y"),
            },
            "warning",
            f"duplicate of row {i + 2}",
        )
    return rows


def workbook_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for values, *_ in rows:
        ws.append([values[k] for k in KEYS])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # keep "=..." as text, as a user typing it with a quote would
        if isinstance(values["internal_date"], datetime):
            ws.cell(row=ws.max_row, column=5).number_format = "dd-mm-yyyy"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def upload(client, path, content, name="schedule.xlsx", **fields):
    file = SimpleUploadedFile(name, content, content_type=XLSX)
    return client.post(f"/api/v1/uploads/{path}", {"file": file, **fields}, format="multipart")


def test_the_generator_really_makes_200_rows():
    rows = build_200_rows()
    assert len(rows) == 200
    from collections import Counter

    assert Counter(status for _, status, _ in rows) == {"ok": 131, "error": 40, "warning": 29}


def test_200_row_file_with_deliberate_errors_is_validated_row_by_row(login, users, world):
    rows = build_200_rows()
    response = upload(login(users["a_planner"]), "preview", workbook_bytes(rows), mode="insert")
    assert response.status_code == 200, response.content
    body = response.json()

    assert body["summary"] == {
        "total": 200, "ok": 131, "warning": 29, "error": 40,
        "insert": 160, "update": 0, "skip": 0, "unchanged": 0,
    }  # fmt: skip
    assert len(body["rows"]) == 200

    for result, (values, expected, note) in zip(body["rows"], rows, strict=True):
        position = result["row"] - 2
        assert rows[position] == (values, expected, note)  # row numbers follow Excel (header is 1)
        assert result["status"] == expected, f"row {result['row']} ({note}): {result['messages']}"
        if expected == "error":
            assert result["action"] == "error" and result["data"] is None
            assert any(m["level"] == "error" for m in result["messages"])
        else:
            assert result["action"] == "insert" and result["data"]["dispatch_qty"] > 0


def test_every_date_and_quantity_style_is_read_correctly(login, users, world):
    rows = build_200_rows()
    body = upload(login(users["a_planner"]), "preview", workbook_bytes(rows), mode="insert").json()
    for i in range(120):
        c, p, d = COMBOS[i]
        data = body["rows"][i]["data"]
        assert data["internal_dispatch_date"] == d.isoformat(), f"row {i + 2}"
        assert data["dispatch_qty"] == (
            500 + i if i % 4 in (0, 1) else 1000 + i if i % 4 == 2 else 700 + i
        )


def test_messages_name_the_problem_precisely(login, users, world):
    rows = build_200_rows()
    by_note = {}
    body = upload(login(users["a_planner"]), "preview", workbook_bytes(rows), mode="insert").json()
    for result, (_, _, note) in zip(body["rows"], rows, strict=True):
        by_note.setdefault(note, []).append(result)

    assert "NOPE0" in by_note["unknown code"][0]["messages"][0]["text"]
    assert "future" in by_note["future actual"][0]["messages"][0]["text"]
    assert by_note["actual qty without date"][0]["messages"][0]["field"] == "actual_qty"
    assert {m["field"] for m in by_note["three problems"][0]["messages"]} == {
        "part_number",
        "dispatch_qty",
        "internal_date",
    }
    assert any("not linked" in m["text"] for m in by_note["not linked"][0]["messages"])
    assert "row 2" in by_note["duplicate of row 2"][0]["messages"][0]["text"]
    assert "matched by name" in by_note["code unknown"][0]["messages"][0]["text"]


def test_preview_never_saves_anything(login, users, world):
    upload(login(users["a_planner"]), "preview", workbook_bytes(build_200_rows()), mode="insert")
    assert ScheduleLine.objects.count() == 0 and UploadBatch.objects.count() == 0


def test_commit_with_errors_is_refused_and_saves_nothing(login, users, world):
    response = upload(
        login(users["a_planner"]), "commit", workbook_bytes(build_200_rows()), mode="insert"
    )
    assert response.status_code == 400 and "40 row(s) have errors" in response.json()["file"]
    assert ScheduleLine.objects.count() == 0 and UploadBatch.objects.count() == 0  # all or nothing


def test_commit_skipping_error_rows_loads_the_rest_in_one_batch(login, users, world):
    client = login(users["a_planner"])
    response = upload(
        client, "commit", workbook_bytes(build_200_rows()), mode="insert", skip_errors="true"
    )
    assert response.status_code == 201, response.content
    body = response.json()
    assert (
        body["rows_total"],
        body["rows_inserted"],
        body["rows_error"],
        body["rows_warning"],
    ) == (200, 160, 40, 29)
    assert ScheduleLine.objects.count() == 160
    assert ScheduleLine.objects.filter(upload_batch_id=body["id"]).count() == 160
    partial = ScheduleLine.objects.filter(status="dispatched")
    assert partial.count() == 11 and all(line.actual_dispatched_qty == 250 for line in partial)
    assert (
        AuditLog.objects.filter(model_label="scheduler.ScheduleLine", action="create").count()
        == 160
    )

    history = client.get("/api/v1/uploads/history").json()["results"]
    assert [(h["file_name"], h["rows_inserted"], h["username"]) for h in history] == [
        ("schedule.xlsx", 160, "a_planner")
    ]

    report = client.get(f"/api/v1/uploads/history/{body['id']}/errors.xlsx")
    sheet = load_workbook(io.BytesIO(report.content)).active
    assert report.status_code == 200 and sheet.max_row == 1 + 40 + 29
    assert [c.value for c in sheet[1]][:3] == ["Row", "Result", "Problem"]


def test_uploading_the_same_file_again_inserts_nothing_in_insert_mode(login, users, world):
    client = login(users["a_planner"])
    content = workbook_bytes(build_200_rows())
    upload(client, "commit", content, mode="insert", skip_errors="true")
    again = upload(client, "preview", content, mode="insert").json()["summary"]
    # Every clean row is now a duplicate of a saved line and is skipped; nothing new to insert.
    assert again["insert"] == 0 and again["skip"] > 0
    refused = upload(client, "commit", content, mode="insert", skip_errors="true")
    assert refused.status_code == 400 and "nothing to save" in refused.json()["file"]
    assert ScheduleLine.objects.count() == 160


def test_update_mode_changes_matching_lines_and_keeps_recorded_actuals(
    login, users, world, company_a
):
    client = login(users["a_planner"])
    line = ScheduleLine.objects.create(
        company=company_a, customer=world[0]["C1"], part=world[1]["P1"], dispatch_qty=100,
        internal_dispatch_date=date(2026, 11, 5), actual_dispatch_date=date(2026, 11, 3),
        actual_dispatched_qty=60, status="dispatched", remarks="kept",
    )  # fmt: skip
    rows = [
        (
            {
                "customer_code": "C1",
                "part_number": "P1",
                "dispatch_qty": 250,
                "internal_date": "05-11-2026",
                "actual_date": "",
                "actual_qty": "",
                "remarks": "",
                "customer_name": "",
            },
            "ok",
            "",
        )
    ]
    summary = upload(client, "preview", workbook_bytes(rows), mode="upsert").json()["summary"]
    assert (summary["update"], summary["insert"]) == (1, 0)

    assert upload(client, "commit", workbook_bytes(rows), mode="upsert").status_code == 201
    line.refresh_from_db()
    assert line.dispatch_qty == 250  # changed
    assert (line.actual_dispatch_date, line.actual_dispatched_qty, line.remarks) == (
        date(2026, 11, 3),
        60,
        "kept",
    )
    assert line.status == "dispatched" and ScheduleLine.objects.count() == 1


def test_update_mode_blocks_duplicate_rows_in_the_file(login, users, world):
    row = {
        "customer_code": "C1",
        "part_number": "P1",
        "dispatch_qty": 5,
        "internal_date": "05-11-2026",
        "customer_name": "",
        "actual_date": "",
        "actual_qty": "",
        "remarks": "",
    }
    body = upload(
        login(users["a_planner"]),
        "preview",
        workbook_bytes([(row, "", ""), (row, "", "")]),
        mode="upsert",
    ).json()
    assert [r["status"] for r in body["rows"]] == ["ok", "error"]


def test_error_report_downloads_as_excel_and_never_runs_formulas(login, users, world):
    evil = {
        "customer_code": "C1",
        "customer_name": "",
        "part_number": '=HYPERLINK("http://evil")',
        "dispatch_qty": 5,
        "internal_date": "05-11-2026",
        "actual_date": "",
        "actual_qty": "",
        "remarks": "=1+1",
    }
    response = upload(
        login(users["a_planner"]), "error-report", workbook_bytes([(evil, "", "")]), mode="insert"
    )
    assert response.status_code == 200 and response["Content-Type"] == XLSX
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2
    cell = sheet.cell(row=2, column=6)  # the part number as uploaded
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    assert sheet.cell(row=2, column=2).value == "Error"


def test_template_has_the_eight_columns_and_instructions(login, users):
    response = login(users["a_viewer"]).get("/api/v1/uploads/template")
    wb = load_workbook(io.BytesIO(response.content))
    assert response.status_code == 200 and wb.sheetnames == ["Schedule", "How to fill"]
    assert [c.value for c in wb["Schedule"][1]] == HEADERS


def test_the_template_round_trips_through_the_importer(login, users, world):
    client = login(users["a_planner"])
    wb = load_workbook(io.BytesIO(client.get("/api/v1/uploads/template").content))
    ws = wb["Schedule"]
    ws.append(["C1", "", "P1", 10, "05-11-2026"])
    buffer = io.BytesIO()
    wb.save(buffer)
    body = upload(client, "preview", buffer.getvalue(), mode="insert").json()
    assert body["summary"]["insert"] == 1


def test_csv_files_are_accepted(login, users, world):
    csv_text = "Customer Code,Part Number,Dispatch Qty,Internal Dispatch Date\nC1,P1,100,05/11/2026\nC9,P1,100,05/11/2026\n"
    body = upload(
        login(users["a_planner"]),
        "preview",
        csv_text.encode("utf-8-sig"),
        name="s.csv",
        mode="insert",
    ).json()
    assert [r["status"] for r in body["rows"]] == ["ok", "error"]


@pytest.mark.parametrize(
    "name, content, text",
    [
        ("notes.txt", b"hello", "xlsx or .csv"),
        ("bad.xlsx", b"not excel", "valid .xlsx"),
        ("empty.csv", b"", "empty"),
    ],
)
def test_unusable_files_get_a_plain_message(login, users, world, name, content, text):
    response = upload(login(users["a_planner"]), "preview", content, name=name, mode="insert")
    assert response.status_code == 400 and text in response.json()["file"]


def test_missing_file_and_bad_mode(login, users, world):
    client = login(users["a_planner"])
    assert client.post("/api/v1/uploads/preview", {}, format="multipart").status_code == 400
    assert (
        upload(client, "preview", workbook_bytes(build_200_rows()[:1]), mode="nonsense").status_code
        == 400
    )


def test_uploads_are_planner_and_admin_only_but_template_is_open(login, users, world):
    content = workbook_bytes(build_200_rows()[:3])
    assert upload(login(users["a_viewer"]), "preview", content, mode="insert").status_code == 403
    assert upload(login(users["a_viewer"]), "commit", content, mode="insert").status_code == 403
    assert login(users["a_viewer"]).get("/api/v1/uploads/history").status_code == 403
    assert upload(login(users["a_admin"]), "preview", content, mode="insert").status_code == 200


def test_company_a_cannot_match_company_bs_customers_or_parts(login, users, world, company_b):
    other = Customer.objects.create(company=company_b, code="B-ONLY", name="B only")
    part = Part.objects.create(
        company=company_b,
        part_number="B-PART",
        name="x",
        casting_weight_kg=1,
        pouring_weight_per_mould_kg=2,
    )
    PartCustomer.objects.create(company=company_b, customer=other, part=part, is_default=True)
    row = {
        "customer_code": "B-ONLY",
        "customer_name": "",
        "part_number": "B-PART",
        "dispatch_qty": 5,
        "internal_date": "05-11-2026",
        "actual_date": "",
        "actual_qty": "",
        "remarks": "",
    }
    body = upload(
        login(users["a_planner"]), "preview", workbook_bytes([(row, "", "")]), mode="insert"
    ).json()
    assert body["rows"][0]["status"] == "error"
    assert {m["field"] for m in body["rows"][0]["messages"]} == {"customer_code", "part_number"}


def test_batch_history_and_error_report_are_per_company(login, users, world, company_a, company_b):
    batch = UploadBatch.objects.create(company=company_b, file_name="b.xlsx", mode="insert")
    client = login(users["a_planner"])
    assert client.get("/api/v1/uploads/history").json()["results"] == []
    assert client.get(f"/api/v1/uploads/history/{batch.pk}/errors.xlsx").status_code == 404


def test_a_large_file_is_validated_quickly(login, users, world):
    rows = []
    for i in range(2000):
        c, p, d = COMBOS[i % len(COMBOS)]
        rows.append(
            (
                {
                    "customer_code": c,
                    "customer_name": "",
                    "part_number": p,
                    "dispatch_qty": 10 + i,
                    "internal_date": d.strftime("%d-%m-%Y"),
                    "actual_date": "",
                    "actual_qty": "",
                    "remarks": "",
                },
                "",
                "",
            )
        )
    content = workbook_bytes(rows)
    started = time.perf_counter()
    response = upload(login(users["a_planner"]), "preview", content, mode="insert")
    elapsed = time.perf_counter() - started
    assert response.status_code == 200 and response.json()["summary"]["total"] == 2000
    assert elapsed < 5, f"validating 2,000 rows took {elapsed:.1f}s"
