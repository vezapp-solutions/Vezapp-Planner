"""Excel files the scheduler hands back: the upload template and the error report."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .importer import COLUMNS, RowResult

HEADER_FILL = PatternFill("solid", fgColor="0B6F67")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", fgColor="F8E4E1")
AMBER = PatternFill("solid", fgColor="F6EAD0")


def put(sheet, row: int, column: int, value):
    """Write a cell as plain text so a value starting with = + - @ can never run as a formula."""
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def template_workbook() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Schedule"
    for column, (_, header) in enumerate(COLUMNS, start=1):
        cell = put(sheet, 1, column, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    sheet.freeze_panes = "A2"
    _widths(sheet, [16, 30, 16, 12, 22, 22, 22, 36])
    for row in range(2, 502):
        sheet.cell(row=row, column=5).number_format = "dd-mm-yyyy"
        sheet.cell(row=row, column=6).number_format = "dd-mm-yyyy"

    help_sheet = wb.create_sheet("How to fill")
    lines = [
        "Fill one row per dispatch requirement on the Schedule sheet. Keep the header row as it is.",
        "",
        "Customer Code or Customer Name: at least one. The code is tried first, then the exact name.",
        "Part Number: must exist in the Part Master.",
        "Dispatch Qty: a whole number of pieces, 1 or more.",
        "Internal Dispatch Date: your plant's target date. Use DD-MM-YYYY (for example 05-11-2026).",
        "Actual Dispatch Date: leave blank until shipped. It cannot be in the future.",
        "Actual Dispatched Qty: optional, for a partial dispatch. Needs an actual date.",
        "Remarks: optional, up to 500 characters.",
        "",
        "Dates may be typed as DD-MM-YYYY or DD/MM/YYYY, or be real Excel dates.",
        "Both .xlsx and .csv files are accepted, up to 5,000 rows and 5 MB.",
        "You always see a preview first. Nothing is saved until you confirm.",
    ]
    for index, text in enumerate(lines, start=1):
        put(help_sheet, index, 1, text)
    help_sheet.column_dimensions["A"].width = 110
    help_sheet["A1"].font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def error_report(results: list[RowResult]) -> bytes:
    """Every row that has an error or a warning, with what is wrong and the values as uploaded."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Problems"
    headers = ["Row", "Result", "Problem", *[h for _, h in COLUMNS]]
    for column, header in enumerate(headers, start=1):
        cell = put(sheet, 1, column, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    sheet.freeze_panes = "A2"

    out_row = 2
    for result in results:
        if result.status == "ok":
            continue
        fill = RED if result.status == "error" else AMBER
        values = [
            result.row_number,
            "Error" if result.status == "error" else "Warning",
            "\n".join(f"{m.field.replace('_', ' ')}: {m.text}" for m in result.messages),
            *[result.raw.get(key, "") for key, _ in COLUMNS],
        ]
        for column, value in enumerate(values, start=1):
            cell = put(sheet, out_row, column, value)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        out_row += 1
    _widths(sheet, [8, 10, 60, 16, 28, 16, 12, 20, 20, 16, 30])
    if out_row == 2:
        put(sheet, 2, 1, "No problems found.")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_from_stored(rows: list[dict]) -> bytes:
    """Rebuild the error report of a committed upload from the rows kept on its history record."""
    from .importer import Message

    results = [
        RowResult(
            row_number=r["row"],
            status=r["status"],
            action=r["action"],
            messages=[Message(m["level"], m["field"], m["text"]) for m in r["messages"]],
            raw=r["raw"],
        )
        for r in rows
    ]
    return error_report(results)
