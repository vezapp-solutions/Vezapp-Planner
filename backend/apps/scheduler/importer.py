"""Reading and validating a monthly schedule file (.xlsx or .csv).

Pure Python: no Django and no database. The caller supplies the lookups (customers, parts,
existing lines) and today's date, which keeps every rule unit-testable.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# key, header shown in the template
COLUMNS = [
    ("customer_code", "Customer Code"),
    ("customer_name", "Customer Name"),
    ("part_number", "Part Number"),
    ("dispatch_qty", "Dispatch Qty"),
    ("internal_date", "Internal Dispatch Date"),
    ("actual_date", "Actual Dispatch Date"),
    ("actual_qty", "Actual Dispatched Qty"),
    ("remarks", "Remarks"),
]
HEADER_OF = dict(COLUMNS)

SYNONYMS = {
    "customercode": "customer_code",
    "custcode": "customer_code",
    "customername": "customer_name",
    "customer": "customer_name",
    "partnumber": "part_number",
    "partno": "part_number",
    "partnum": "part_number",
    "part": "part_number",
    "dispatchqty": "dispatch_qty",
    "dispatchquantity": "dispatch_qty",
    "qty": "dispatch_qty",
    "quantity": "dispatch_qty",
    "internaldispatchdate": "internal_date",
    "internaldate": "internal_date",
    "dispatchdate": "internal_date",
    "actualdispatchdate": "actual_date",
    "actualdate": "actual_date",
    "actualdispatchedqty": "actual_qty",
    "actualdispatchqty": "actual_qty",
    "actualqty": "actual_qty",
    "remarks": "remarks",
    "remark": "remarks",
    "notes": "remarks",
}

MIN_SERIAL, MAX_SERIAL = 18264, 73415  # 1950-01-01 to 2100-12-31 as Excel serial numbers
MIN_YEAR, MAX_YEAR = 2000, 2100

Level = Literal["error", "warning"]
Status = Literal["ok", "warning", "error"]
Action = Literal["insert", "update", "skip", "unchanged", "error"]
Mode = Literal["insert", "upsert"]


class FileProblem(Exception):
    """The whole file cannot be used (wrong type, too big, missing columns)."""


@dataclass
class RawRow:
    number: int  # the row number as the user sees it in Excel (header is row 1)
    values: dict[str, Any]


@dataclass(frozen=True)
class CustomerInfo:
    id: int
    code: str
    name: str
    active: bool = True


@dataclass(frozen=True)
class PartInfo:
    id: int
    number: str
    name: str
    active: bool = True


@dataclass(frozen=True)
class ExistingLine:
    id: int
    qty: int
    actual_date: date | None
    actual_qty: int | None
    remarks: str
    status: str


@dataclass
class Lookups:
    customers_by_code: dict[str, CustomerInfo] = field(default_factory=dict)
    customers_by_name: dict[str, CustomerInfo] = field(default_factory=dict)
    parts_by_number: dict[str, PartInfo] = field(default_factory=dict)
    links: set[tuple[int, int]] = field(default_factory=set)  # (part_id, customer_id)
    existing: dict[tuple[int, int, date], list[ExistingLine]] = field(default_factory=dict)


@dataclass
class Message:
    level: Level
    field: str
    text: str

    def as_dict(self) -> dict[str, str]:
        return {"level": self.level, "field": self.field, "text": self.text}


@dataclass
class RowResult:
    row_number: int
    status: Status
    action: Action
    messages: list[Message]
    raw: dict[str, str]
    data: dict[str, Any] | None = None
    existing_id: int | None = None

    def as_dict(self) -> dict[str, Any]:
        data = None
        if self.data:
            data = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in self.data.items()}
        return {
            "row": self.row_number,
            "status": self.status,
            "action": self.action,
            "messages": [m.as_dict() for m in self.messages],
            "raw": self.raw,
            "data": data,
            "existing_id": self.existing_id,
        }


# --- reading the file ----------------------------------------------------------------------------
def _norm_header(text: Any) -> str:
    return re.sub(r"[^a-z]", "", str(text or "").lower())


def _map_headers(headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, header in enumerate(headers):
        key = SYNONYMS.get(_norm_header(header))
        if key and key not in mapping.values():
            mapping[index] = key
    return mapping


def _check_columns(mapping: dict[int, str]) -> None:
    have = set(mapping.values())
    missing = []
    if "part_number" not in have:
        missing.append("Part Number")
    if "dispatch_qty" not in have:
        missing.append("Dispatch Qty")
    if "internal_date" not in have:
        missing.append("Internal Dispatch Date")
    if not have & {"customer_code", "customer_name"}:
        missing.append("Customer Code or Customer Name")
    if missing:
        raise FileProblem(
            "The first row must be a header row. Missing column(s): "
            + ", ".join(missing)
            + ". Use “Download template” for the exact layout."
        )


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _rows_from_matrix(matrix) -> list[RawRow]:
    iterator = iter(matrix)
    try:
        headers = list(next(iterator))
    except StopIteration as exc:
        raise FileProblem("The file is empty.") from exc
    mapping = _map_headers(headers)
    _check_columns(mapping)

    rows: list[RawRow] = []
    for number, cells in enumerate(iterator, start=2):
        values = {key: (cells[i] if i < len(cells) else None) for i, key in mapping.items()}
        if all(_is_blank(v) for v in values.values()):
            continue  # a fully blank row is just spacing
        rows.append(RawRow(number, values))
        if len(rows) > MAX_ROWS:
            raise FileProblem(f"The file has more than {MAX_ROWS:,} rows. Split it by month.")
    if not rows:
        raise FileProblem("The file has a header but no data rows.")
    return rows


def read_table(filename: str, content: bytes) -> list[RawRow]:
    """Rows of an .xlsx or .csv upload. Raises FileProblem when the file cannot be used."""
    if len(content) > MAX_BYTES:
        raise FileProblem("The file is larger than 5 MB.")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise FileProblem("This is not a valid .xlsx file.")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.worksheets[0]
            matrix = [list(r) for r in sheet.iter_rows(values_only=True)]
        except Exception as exc:  # noqa: BLE001 - any parser failure means an unusable file
            raise FileProblem("The Excel file could not be read. Re-save it as .xlsx.") from exc
        return _rows_from_matrix(matrix)
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return _rows_from_matrix(list(csv.reader(io.StringIO(text), delimiter=delimiter)))
    raise FileProblem("Upload an .xlsx or .csv file.")


# --- parsing single cells --------------------------------------------------------------------------
_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d")


def parse_date(value: Any) -> date:
    """DD-MM-YYYY, DD/MM/YYYY, ISO, a real Excel date, or an Excel serial number."""
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    elif isinstance(value, bool):
        raise ValueError("Not a date.")
    elif isinstance(value, (int, float, Decimal)) or (
        isinstance(value, str) and re.fullmatch(r"\d{5}(\.\d+)?", value.strip())
    ):
        serial = float(value)
        if not (MIN_SERIAL <= serial <= MAX_SERIAL):
            raise ValueError(f"{value} is not a date.")
        result = from_excel(int(serial)).date()
    elif isinstance(value, str):
        text = value.strip().replace(" ", "")
        for fmt in _DATE_FORMATS:
            try:
                result = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"“{value.strip()}” is not a date. Use DD-MM-YYYY.")
    else:
        raise ValueError("Not a date.")
    if not (MIN_YEAR <= result.year <= MAX_YEAR):
        raise ValueError(f"{result.strftime('%d-%m-%Y')} is outside the supported years.")
    return result


def parse_qty(value: Any) -> int:
    """A whole number of pieces, 1 or more. Accepts 600, 600.0, '1,200' and '1 200'."""
    if isinstance(value, bool):
        raise ValueError("Not a quantity.")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        text = str(value).strip().replace(",", "").replace(" ", "").replace(" ", "")
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(f"“{str(value).strip()}” is not a number.") from exc
    if not number.is_finite():
        raise ValueError("Not a quantity.")
    if number != number.to_integral_value():
        raise ValueError(f"{number.normalize():f} is not a whole number of pieces.")
    if number < 1:
        raise ValueError("The quantity must be at least 1.")
    if number > 10_000_000:
        raise ValueError("That quantity is unrealistically large.")
    return int(number)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 1001.0 typed as a code should read 1001
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    return str(value).strip()


def _key(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


# --- validating rows ----------------------------------------------------------------------------------
def validate_rows(rows: list[RawRow], lookups: Lookups, today: date, mode: Mode) -> list[RowResult]:
    results: list[RowResult] = []
    first_seen: dict[tuple[int, int, date], int] = {}

    for row in rows:
        v = row.values
        raw = {key: _text(v.get(key)) for key, _ in COLUMNS}
        messages: list[Message] = []

        def error(field_name: str, text: str, messages=messages) -> None:
            messages.append(Message("error", field_name, text))

        def warn(field_name: str, text: str, messages=messages) -> None:
            messages.append(Message("warning", field_name, text))

        # customer: code first, then exact name
        customer = None
        code, cname = _key(raw["customer_code"]), _key(raw["customer_name"])
        if code:
            customer = lookups.customers_by_code.get(code)
            if customer is None and not cname:
                error("customer_code", f"Customer code “{raw['customer_code']}” does not exist.")
        if customer is None and cname:
            customer = lookups.customers_by_name.get(cname)
            if customer is None:
                error("customer_name", f"Customer “{raw['customer_name']}” does not exist.")
            elif code:
                warn(
                    "customer_code",
                    f"Code “{raw['customer_code']}” is unknown; matched by name to {customer.code}.",
                )
        elif customer is not None and cname and _key(customer.name) != cname:
            warn(
                "customer_name",
                f"The name given differs from the name of customer {customer.code} ({customer.name}).",
            )
        if not code and not cname:
            error("customer_code", "Give a customer code or name.")
        if customer is not None and not customer.active:
            warn("customer_code", f"Customer {customer.code} is marked inactive.")

        # part
        part = None
        if not raw["part_number"]:
            error("part_number", "Part number is missing.")
        else:
            part = lookups.parts_by_number.get(_key(raw["part_number"]))
            if part is None:
                error("part_number", f"Part “{raw['part_number']}” does not exist.")
            elif not part.active:
                warn("part_number", f"Part {part.number} is marked inactive.")
        if (
            customer is not None
            and part is not None
            and (part.id, customer.id) not in lookups.links
        ):
            warn(
                "part_number",
                f"Part {part.number} is not linked to customer {customer.code} in the Part Master.",
            )

        # quantity
        qty = None
        if _is_blank(v.get("dispatch_qty")):
            error("dispatch_qty", "Dispatch quantity is missing.")
        else:
            try:
                qty = parse_qty(v["dispatch_qty"])
            except ValueError as exc:
                error("dispatch_qty", str(exc))

        # dates
        internal = None
        if _is_blank(v.get("internal_date")):
            error("internal_date", "Internal dispatch date is missing.")
        else:
            try:
                internal = parse_date(v["internal_date"])
            except ValueError as exc:
                error("internal_date", str(exc))

        actual = None
        if not _is_blank(v.get("actual_date")):
            try:
                actual = parse_date(v["actual_date"])
                if actual > today:
                    error(
                        "actual_date",
                        f"The actual dispatch date {actual.strftime('%d-%m-%Y')} is in the future.",
                    )
            except ValueError as exc:
                error("actual_date", str(exc))

        actual_qty = None
        if not _is_blank(v.get("actual_qty")):
            try:
                actual_qty = parse_qty(v["actual_qty"])
            except ValueError as exc:
                error("actual_qty", str(exc))
            else:
                if _is_blank(v.get("actual_date")):
                    error(
                        "actual_qty", "An actual dispatched quantity needs an actual dispatch date."
                    )

        remarks = raw["remarks"]
        if len(remarks) > 500:
            error("remarks", "Remarks are longer than 500 characters.")

        # duplicates and matching existing lines
        action: Action = "insert"
        existing_id = None
        if customer is not None and part is not None and internal is not None:
            key = (customer.id, part.id, internal)
            earlier = first_seen.get(key)
            if earlier is None:
                first_seen[key] = row.number
            elif mode == "upsert":
                error(
                    "internal_date",
                    f"Same customer, part and date as row {earlier}. Combine them into one row.",
                )
            else:
                warn("internal_date", f"Same customer, part and date as row {earlier}.")

            matches = lookups.existing.get(key, [])
            if len(matches) > 1 and mode == "upsert":
                error(
                    "internal_date",
                    f"Matches {len(matches)} lines already in the schedule. Edit them in the grid.",
                )
            elif matches:
                existing_id = matches[0].id
                if mode == "insert":
                    action = "skip"
                    warn("internal_date", "Already in the schedule, so this row is skipped.")
                else:
                    action = "update"

        has_error = any(m.level == "error" for m in messages)
        data = None
        if not has_error:
            data = {
                "customer_id": customer.id,
                "part_id": part.id,
                "dispatch_qty": qty,
                "internal_dispatch_date": internal,
                "actual_dispatch_date": actual,
                "actual_dispatched_qty": actual_qty,
                "remarks": remarks,
            }
            if action == "update":
                old = next(m for m in lookups.existing[(customer.id, part.id, internal)])
                changed = (
                    qty != old.qty
                    or (actual is not None and actual != old.actual_date)
                    or (actual_qty is not None and actual_qty != old.actual_qty)
                    or (remarks != "" and remarks != old.remarks)
                )
                if not changed:
                    action = "unchanged"
        else:
            action = "error"

        status: Status = "error" if has_error else ("warning" if messages else "ok")
        results.append(RowResult(row.number, status, action, messages, raw, data, existing_id))
    return results


def summarize(results: list[RowResult]) -> dict[str, int]:
    out = {
        "total": len(results),
        "ok": 0,
        "warning": 0,
        "error": 0,
        "insert": 0,
        "update": 0,
        "skip": 0,
        "unchanged": 0,
    }
    for r in results:
        out[r.status] += 1
        if r.action in ("insert", "update", "skip", "unchanged"):
            out[r.action] += 1
    return out


def month_end(day: date) -> date:
    first_next = (day.replace(day=28) + timedelta(days=4)).replace(day=1)
    return first_next - timedelta(days=1)
