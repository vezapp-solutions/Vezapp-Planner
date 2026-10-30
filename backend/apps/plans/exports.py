"""Render a Report as PDF (reportlab) or Excel (openpyxl), with the Vezapp logo."""

import io
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.scheduler.reports import put

from .models import Plan
from .reports import Report

LOGO = Path(settings.BRAND_LOGO_PATH)
TEAL = colors.HexColor("#0B6F67")
INK = colors.HexColor("#232625")
LINE = colors.HexColor("#D6D2C9")
ZEBRA = colors.HexColor("#F4F2EE")


def _footer_text(plan: Plan, report: Report) -> str:
    when = (
        plan.calculated_at.astimezone().strftime("%d-%m-%Y %H:%M")
        if plan.calculated_at
        else "not calculated"
    )
    return f"{settings.APP_NAME} · {plan.company.name} · plan v{plan.version} ({plan.get_status_display().lower()}), calculated {when}"


# --- Excel ---------------------------------------------------------------------------------------------
def to_xlsx(plan: Plan, report: Report) -> bytes:
    return render_xlsx(report, _footer_text(plan, report))


def render_xlsx(report: Report, footer: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    row = 1
    if LOGO.exists():
        image = XlImage(str(LOGO))
        image.height, image.width = 40, int(40 * image.width / image.height)
        ws.add_image(image, "A1")
        row = 4
    put(ws, row, 1, report.title).font = Font(bold=True, size=14)
    put(ws, row + 1, 1, report.subtitle)
    put(ws, row + 2, 1, footer).font = Font(color="666666", size=9)
    row += 4
    widest = 0
    for table in report.tables:
        put(ws, row, 1, table.title).font = Font(bold=True, size=12)
        row += 1
        for c, name in enumerate(table.columns, start=1):
            cell = put(ws, row, c, name)
            cell.font, cell.fill = (
                Font(bold=True, color="FFFFFF"),
                PatternFill("solid", fgColor="0B6F67"),
            )
        row += 1
        for values in table.rows:
            for c, value in enumerate(values, start=1):
                put(ws, row, c, value).alignment = Alignment(vertical="top")
            row += 1
        widest = max(widest, len(table.columns))
        row += 1
    for c in range(1, widest + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20 if c > 1 else 24
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDF -----------------------------------------------------------------------------------------------
def to_pdf(plan: Plan, report: Report) -> bytes:
    return render_pdf(report, _footer_text(plan, report))


def render_pdf(report: Report, footer: str) -> bytes:
    buffer = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=26 * mm, bottomMargin=16 * mm,
                            title=report.title, author=settings.APP_NAME)  # fmt: skip
    styles = getSampleStyleSheet()
    h2 = ParagraphStyle("h2", parent=styles["Heading3"], textColor=INK, spaceBefore=8, spaceAfter=4)
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10, textColor=INK)

    def chrome(canvas, document):
        canvas.saveState()
        width, height = page
        if LOGO.exists():
            from reportlab.lib.utils import ImageReader

            reader = ImageReader(str(LOGO))
            iw, ih = reader.getSize()
            h = 14 * mm
            canvas.drawImage(
                reader, 12 * mm, height - 20 * mm, width=h * iw / ih, height=h, mask="auto"
            )
        canvas.setFont("Helvetica-Bold", 14)
        canvas.setFillColor(INK)
        canvas.drawRightString(width - 12 * mm, height - 12 * mm, report.title)
        canvas.setFont("Helvetica", 9)
        canvas.drawRightString(width - 12 * mm, height - 17 * mm, report.subtitle)
        canvas.setStrokeColor(TEAL)
        canvas.setLineWidth(1.2)
        canvas.line(12 * mm, height - 22 * mm, width - 12 * mm, height - 22 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#565A57"))
        canvas.drawString(12 * mm, 8 * mm, footer)
        canvas.drawRightString(width - 12 * mm, 8 * mm, f"Page {document.page}")
        canvas.restoreState()

    story = []
    usable = page[0] - 24 * mm
    for t in report.tables:
        story.append(Paragraph(t.title, h2))
        head = [
            Paragraph(f"<b>{c}</b>", ParagraphStyle("hd", parent=cell, textColor=colors.white))
            for c in t.columns
        ]
        body = [[Paragraph(str(v), cell) for v in row] for row in t.rows]
        weights = t.widths or [1] * len(t.columns)
        total = sum(weights)
        grid = Table([head, *body], colWidths=[usable * w / total for w in weights], repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), TEAL),
            ("GRID", (0, 0), (-1, -1), 0.25, LINE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        style += [("BACKGROUND", (0, i), (-1, i), ZEBRA) for i in range(2, len(body) + 1, 2)]
        grid.setStyle(TableStyle(style))
        story += [grid, Spacer(1, 4 * mm)]
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
    return buffer.getvalue()
